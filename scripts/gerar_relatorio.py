"""Relatório acadêmico derivado dos registros; executar com o Python de artefatos."""
from pathlib import Path
import hashlib
import json
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"outputs"
analysis=json.loads((OUT/"analysis/summary.json").read_text(encoding="utf-8"))
summaries=analysis["summaries"]
selected=analysis["selected"]
eval_path=OUT/"llm/evaluation/summary.json"
llm=json.loads(eval_path.read_text(encoding="utf-8")) if eval_path.exists() else None
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.left_margin=s.right_margin=s.top_margin=s.bottom_margin=Inches(1)
s.header_distance=s.footer_distance=Inches(.5)
s.different_first_page_header_footer=True
normal=doc.styles["Normal"]
normal.font.name="Calibri";normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.15
normal.paragraph_format.space_after=Pt(8)
for name,size in [("Heading 1",16),("Heading 2",13)]:
    st=doc.styles[name];st.font.name="Calibri";st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string("2E74B5")
    st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(10);st.paragraph_format.keep_with_next=True
title=doc.styles["Title"];title.font.name="Calibri";title.font.size=Pt(30);title.font.color.rgb=RGBColor(0,0,0)
title.paragraph_format.line_spacing=1.1
for border in list(title._element.iter(qn("w:pBdr"))):
    border.getparent().remove(border)
cap=doc.styles["Caption"];cap.font.name="Calibri";cap.font.size=Pt(9);cap.font.color.rgb=RGBColor.from_string("586575");cap.font.italic=True
header=s.header.paragraphs[0];header.text="Tech Challenge Fase 2\tIADT • FIAP"
header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_ALIGN_PARAGRAPH.RIGHT)
for r in header.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string("637181")
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Página ")
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
for r in footer.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string("637181")
required_text=[]
figure_number=0


def fmt(value,digits=3):
    if value is None:return "—"
    result=f"{value:.{digits}f}".replace(".",",")
    required_text.append(result)
    return result


def p(text):return doc.add_paragraph(text)
def h(text):return doc.add_heading(text,level=2)
def page(title):doc.add_page_break();doc.add_heading(title,level=1)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";t.autofit=False
    widths=widths or [6.5/len(headers)]*len(headers)
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for c,text,w in zip(t.rows[0].cells,headers,widths):
        c.text=text;c.width=Inches(w)
        shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"F1F4F8");c._tc.get_or_add_tcPr().append(shade)
    repeat=OxmlElement("w:tblHeader");t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for cell,text,w in zip(cells,row,widths):cell.text=str(text);cell.width=Inches(w)
    for ri,row in enumerate(t.rows):
        keep=OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(keep)
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4);para.paragraph_format.line_spacing=1
                for run in para.runs:
                    run.font.name="Calibri";run.font.size=Pt(9.5)
                    if ri==0:run.bold=True;run.font.color.rgb=RGBColor.from_string("285B86")
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t


def picture(path,caption,height=3.2,width=6.5):
    global figure_number
    source=ROOT/path
    with Image.open(source) as image:w,h=image.size
    factor=min(width/w,height/h)
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(source),width=Inches(w*factor))
    figure_number+=1
    doc.add_paragraph(f"Figura {figure_number}  {caption}",style="Caption")


def result(name,algorithm,config=None):
    return next(r for r in summaries if r["scenario"]==name and r["algorithm"]==algorithm and (config is None or r["configuration"]==config))


banner=p("TECH CHALLENGE FASE 2");banner.alignment=WD_ALIGN_PARAGRAPH.CENTER
banner.paragraph_format.space_before=Pt(85);banner.paragraph_format.space_after=Pt(32)
for run in banner.runs:run.bold=True;run.font.color.rgb=RGBColor.from_string("A67B13");run.font.size=Pt(14)
cover=doc.add_paragraph("Otimização de rotas médicas com algoritmo genético e LLM",style="Title");cover.alignment=WD_ALIGN_PARAGRAPH.CENTER
sub=p("Múltiplos veículos, prioridades e interpretação de soluções verificadas");sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after=Pt(30)
for r in sub.runs:r.font.size=Pt(15);r.font.color.rgb=RGBColor.from_string("35586A")
team=p("Equipe");team.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in team.runs:r.bold=True
for name in ["Lucas dos Santos Batista  RM 376835","Eduardo Fagner Costalonga Ribeiro  RM 376780","Renan Gonçalves Micheletti  RM 376743"]:
    para=p(name);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
date=p("Outubro de 2026");date.alignment=WD_ALIGN_PARAGRAPH.CENTER;date.paragraph_format.space_before=Pt(28)
state=p("Entrega local em preparação  Integração real com Luna pendente" if not llm else "Relatório técnico da entrega local")
state.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in state.runs:r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string("637181")

page("1 Visão geral e escopo")
p("Implementamos um sistema de planejamento da distribuição de medicamentos e insumos por múltiplos veículos, com capacidade, autonomia e prioridades. O núcleo utiliza um algoritmo genético, comparado com uma heurística gulosa e OR-Tools. Uma camada de linguagem organiza fatos de soluções verificadas para produzir instruções, relatórios e respostas a perguntas.")
p("A avaliação mostra que o genético encontra rotas viáveis e melhora alguns casos menores, mas não supera OR-Tools nos benchmarks maiores. A reparação consome uma parte relevante do orçamento, limitando a evolução. Esse resultado orienta a discussão das limitações e evita atribuir à abordagem ganhos que não foram observados.")
p("O cenário hospitalar é fictício e utiliza ruas reais da região central de São Paulo. Os tempos são estimativas sem trânsito em tempo real. Não usamos pacientes, medicamentos reais, estoques ou desfechos clínicos. Os pesos de prioridade representam premissas do experimento.")
table(["Componente","Entrega"],[
    ["Otimização","Pacote Python compartilhado por CLI, notebooks e API"],
    ["Dados","Snapshot OSM, cenários de 12/30/60 entregas, sete dias e três benchmarks"],
    ["Experimentos",f"{analysis['runs']} execuções tabuladas, três configurações e dez sementes finais"],
    ["Interpretação","gpt-6-luna pela Responses API, saída estruturada e orçamento acumulado"],
    ["Documentação","Dois notebooks, gráficos, mapas, relatório e roteiro de apresentação"],
    ["Validação","Casos conhecidos, adulterações, integração HTTP e aceite por evidências"]],[1.35,5.15])
p("A conclusão do projeto é a entrega local validada. A equipe realizará a gravação do vídeo, a publicação e a submissão. O roteiro prevê 14 minutos e 30 segundos, sem incluir essas ações no trabalho automatizado.")
if not llm:p("Nesta versão, a execução real da LLM depende da configuração de uma chave específica. Os testes simulados não substituem esse requisito. A ausência dessa evidência mantém o aceite integral pendente.")

page("2 Modelo de dados e restrições")
p("Cada cenário contém um depósito, entregas com identificadores únicos, carga, prioridade, atendimento e localização, além de uma frota com capacidade e autonomia. A ordem dos nós nas matrizes é depósito seguido pelas entregas. A localização não define a identidade de um pedido: duas entregas no mesmo endereço permanecem distintas.")
table(["Regra","Conferência"],[
    ["Visitas","Todos os pedidos exatamente uma vez, sem IDs desconhecidos"],
    ["Capacidade","Soma da carga inicial de cada rota dentro do limite do veículo"],
    ["Autonomia","Distância completa, incluindo retorno ao depósito"],
    ["Horários","Deslocamento dirigido e atendimento das paradas anteriores"],
    ["Matrizes","Dimensão correta, diagonal zero, valores finitos e pares alcançáveis coerentes"],
    ["Publicação","Rota viável recertificada antes da consulta e da interpretação"]],[1.3,5.2])
p("Os veículos saem em paralelo no horário inicial comum, sem reabastecimento. A carga diminui ao longo das entregas. Não há coletas, janelas de atendimento, restrições térmicas nem incompatibilidade entre tipos de carga. Autonomia significa distância máxima; não representa um modelo energético.")
h("Objetivo entre soluções viáveis")
p("J = distância total em quilômetros + α × soma dos minutos de chegada ponderados pela prioridade. Usamos α = 0,05 km/min. O tempo de atendimento de uma parada afeta a chegada das seguintes, mas não a chegada à própria parada.")
table(["Prioridade","Peso inicial"],[["Regular","1"],["Urgente","3"],["Crítica","10"]],[4.5,2])
p("A avaliação compara primeiro erros estruturais, depois violações de capacidade e autonomia, e só então J. Um percurso inválido não vence por ser mais curto. Nos benchmarks CVRPLIB e no TSP FIAP, usamos exclusivamente a distância original, sem misturar unidades ou prioridades médicas.")

page("3 Adaptação do algoritmo genético")
table(["Aspecto","Base original","Projeto atual"],[
    ["Problema","TSP com um circuito","Múltiplos veículos e restrições"],
    ["Genes","Coordenadas","IDs exclusivos e separadores"],
    ["Seleção","Roleta","Torneio de três"],
    ["Crossover","OX com o mesmo pai no laço","OX com pais distintos quando possível"],
    ["Mutação","Troca entre vizinhos","Troca ou inversão"],
    ["Elitismo","Um indivíduo","Dois indivíduos"],
    ["Verificação","Comprimento do ciclo","Certificação independente de visitas e limites"],
    ["Aplicação","Interface Pygame","Pacote, CLI, API e mapas"]],[1.0,2.35,3.15])
p("O crossover original distingue valores de coordenadas, mas não identifica pedidos. Seu teste de pertencimento pode perder a multiplicidade de entregas no mesmo endereço. O novo OX atua sobre genes únicos e conserva todos os pedidos e separadores.")
p("A representação permite rotas vazias quando um veículo não é usado. A reparação determinística procura realocações e trocas que reduzam violações. Ela preserva as visitas, mas não garante encontrar uma solução para todo problema viável. O código separa avaliação da busca e verificação da solução publicada.")
p("A reprodução da referência remove apenas a dependência gráfica e mantém seus operadores, probabilidade de mutação 0,50, um elite e crossover com o mesmo pai. Essa comparação é restrita ao TSP de 15 pontos compatível com a fonte.")

page("4 Dados e preparação geográfica")
p("Obtivemos uma malha de ruas direcionadas do OpenStreetMap com OSMnx 2.1.1. O recorte vai de -46,69 a -46,61 de longitude e de -23,59 a -23,52 de latitude. Retivemos a maior componente fortemente conexa, com 7.029 nós e 14.036 arestas, permitindo ida e retorno entre os pontos escolhidos.")
p("Cada par usa o caminho de menor tempo estimado na malha. Distância, duração e polilinha são calculadas sobre as mesmas arestas. A velocidade deriva do campo maxspeed quando disponível, de imputação por tipo de rua e de fallback de 30 km/h. Sentidos das ruas são considerados; trânsito atual e restrições de conversão não são modelados.")
table(["Conjunto","Escalas e frota","Sementes"],[
    ["Desenvolvimento","12/2, 30/3 e 60/5 entregas/veículos","617+n"],
    ["Avaliação final","12/2, 30/3 e 60/5","2026+n"],
    ["Semana simulada","Sete dias com 12 entregas e dois veículos","5000 a 5006"],
    ["Variações","Prioridade, capacidade e autonomia","Semente 42 da busca"],
    ["CVRPLIB","A-n32-k5, A-n54-k7 e A-n80-k10","Instâncias públicas originais"]],[1.45,3.35,1.7])
p("Cargas são inteiros de 2 a 10 kg; prioridades são sorteadas com proporções de 55%, 30% e 15%. O atendimento utiliza 180, 300 ou 420 segundos. A capacidade nominal é o teto de 1,35 vezes a carga total dividida pela frota, com autonomia de 50 km. As datas de 5 a 11 de outubro são simulações, sem alegação de operações ocorridas.")
p("As instâncias CVRPLIB preservam os arquivos VRP e SOL, capacidade, frota e arredondamento EUC_2D floor(d+0,5). O endereço original apresentou certificado TLS expirado; a obtenção usou o espelho mantido pelo autor, com commit e hashes registrados e validação TLS ativa. As soluções publicadas foram conferidas independentemente.")
p("A malha, ruas, matrizes e geometrias são incluídas localmente. A atribuição © OpenStreetMap contributors e a licença ODbL 1.0 são conservadas. Uma nova obtenção pode refletir alterações da fonte; os resultados desta entrega usam o snapshot salvo.")

page("5 Experimentos e orçamento de busca")
p("Comparamos três configurações do genético. Crossover, torneio e elitismo permanecem fixos. A seleção da configuração de uso considera apenas a média da razão J do GA/J do guloso nos três cenários de desenvolvimento, com três sementes por configuração. A escolha é registrada antes da avaliação final.")
table(["Configuração","População","Mutação","Crossover"],[["pop80_mut10","80","0,10","0,90"],["pop160_mut20","160","0,20","0,90"],["pop240_mut30","240","0,30","0,90"]],[2.1,1.4,1.5,1.5])
p(f"A configuração selecionada foi {selected}. Cada configuração usa dez sementes nos sete problemas finais: três cenários médicos, três instâncias CVRPLIB e o TSP FIAP. As sementes são 42, 73, 101, 127, 151, 181, 211, 241, 271 e 307. As execuções adicionais cobrem sensibilidade, restrições e a semana simulada.")
p("O teto de 20.000 inclui todas as chamadas de fitness da busca, inclusive inicialização e reparação. Há também um limite de 20.000 tentativas da população, para encerrar alta repetição de indivíduos, e de 120 segundos. Custos de inserção de rotas parciais são contados separadamente. A certificação final fica fora do orçamento de otimização.")
p("Uma primeira implementação contava apenas as tentativas da população. Ao revisar o orçamento, identificamos que as chamadas da reparação não estavam incluídas. Invalidamos o uso desses resultados na comparação final e repetimos os experimentos com a contagem corrigida. A entrega utiliza exclusivamente as execuções com contagem completa.")
p("OR-Tools usa cinco segundos, restrições de capacidade e autonomia e objetivo equivalente, convertido em inteiros com precisão documentada. O guloso escolhe a melhor entre três construções determinísticas. Os orçamentos não são equivalentes; tempo de execução e limite de busca são apresentados separadamente. Os tempos foram medidos em CPU local com outras atividades em paralelo e não constituem um benchmark isolado de hardware.")
p("Cada resultado registra semente, parâmetros, versão dos dados, hash da implementação, duração, fitness, tentativas, diversidade e condição de término. Checkpoints permitem retomar sem repetir uma execução já conferida.")

page("6 Resultados nos cenários médicos")
rows=[]
for n in (12,30,60):
    for algo,label in [("greedy","Guloso"),("ortools","OR-Tools"),("genetic","GA selecionado")]:
        r=result(f"sp_evaluation_{n}",algo,selected if algo=="genetic" else None)
        rows.append([n,label,fmt(r["objective_mean"]),fmt(r["objective_std"]),fmt(r["distance_mean"]/1000),f"{r['feasible']}/{r['runs']}"])
table(["Pedidos","Método","J médio","Desvio de J","km médios","Viáveis"],rows,[.65,1.45,1.05,1.05,1.05,1.25])
picture("outputs/figures/medical_objective.png","Distribuição de J por configuração e comparação com guloso e OR-Tools.",height=2.45)
p("Com 12 entregas, as configurações alcançam o mesmo J de OR-Tools nas sementes finais. Com 30 entregas, o GA selecionado apresenta pequenas melhorias sobre a construção gulosa, mas mantém custo superior ao OR-Tools. No cenário de 60 entregas, a busca preserva a solução inicial e não demonstra ganho sobre o guloso.")
p("A igualdade entre sementes nos problemas maiores não comprova robustez da evolução: pode resultar do orçamento consumido pela reparação. A distância e J devem ser lidos separadamente, pois J também considera chegada ponderada. Todos os resultados viáveis foram certificados quanto às restrições.")

page("7 Convergência e uso do orçamento")
picture("outputs/figures/convergence.png","Melhor J viável observado e interrupções antes de completar gerações.",height=3.2)
p("Nos cenários maiores, a reparação avalia numerosas inserções e trocas. Ao contabilizar esse custo dentro das 20 mil chamadas, parte das execuções termina durante a inicialização. Um ponto isolado no gráfico indica esse término e não deve ser apresentado como convergência de um processo evolutivo completo.")
table(["Condição de término","Execuções GA tabuladas"],[[k,v] for k,v in analysis["termination_counts"].items()],[4.5,2])
p("O limite de tentativas encerra a busca quando a população repete candidatos, mesmo que a quantidade de chamadas novas de fitness seja menor que 20 mil. Esse encerramento evita interpretar cache como processamento adicional. A melhor solução viável já avaliada é conservada quando o orçamento termina em qualquer etapa.")
p("A conclusão é que a reparação implementada funciona como mecanismo de preservação de restrições, mas é cara para este orçamento. Uma reparação incremental, candidatos de inserção mais seletivos e construção inicial com melhor diversidade são propostas de melhoria. Elas exigiriam nova avaliação; não atribuímos ganhos a versões ainda não testadas.")

page("8 Resultados nos benchmarks públicos")
rows=[]
for name in ("A-n32-k5","A-n54-k7","A-n80-k10"):
    for algo,label in [("greedy","Guloso"),("ortools","OR-Tools"),("genetic","GA selecionado")]:
        r=result(name,algo,selected if algo=="genetic" else None)
        rows.append([name,label,fmt(r["distance_mean"],1),fmt(r["gap_mean_percent"],2),f"{r['feasible']}/{r['runs']}"])
table(["Instância","Método","Distância média","Gap médio %","Viáveis"],rows,[1.35,1.5,1.4,1.15,1.1])
picture("outputs/figures/benchmark_gap.png","Distância relativa aos custos publicados 784, 1167 e 1763.",height=2.5)
p("OR-Tools alcançou o custo publicado na instância A-n32-k5 e ficou mais próximo da referência nas demais. O genético permaneceu acima desses custos, com pequenas melhorias sobre o guloso em alguns casos. A-n80-k10 não apresentou ganho do GA sobre a construção inicial neste orçamento.")
p("Os gaps utilizam unidades originais do benchmark. Não são percentuais de quilômetros nem estimativas de atraso hospitalar. As soluções publicadas e todas as soluções experimentais viáveis respeitam a capacidade e o atendimento de cada cliente exatamente uma vez.")

page("9 Comparação com o TSP FIAP")
picture("outputs/figures/fiap_tsp.png","Comprimento dos circuitos no TSP original de 15 pontos.",height=3.35)
reference=result("fiap_tsp_15","fiap_original");new=result("fiap_tsp_15","genetic",selected)
p(f"A referência FIAP apresentou média de {fmt(reference['distance_mean'])} unidades euclidianas. A configuração selecionada apresentou {fmt(new['distance_mean'])}. O guloso e OR-Tools alcançaram o mesmo valor observado pelo novo GA nesse cenário.")
p("Esse resultado mostra vantagem da implementação completa sobre a reprodução da referência, mas não isola o efeito da correção do crossover. Mudaram também a inicialização, a seleção, a mutação, o elitismo e o uso de construção gulosa. Como o novo GA coincide com a solução inicial nesse TSP, não atribuímos a redução exclusivamente à evolução genética.")
p("O comparador mantém a chamada OX com o mesmo pai e a mutação por troca de vizinhos. A nova versão opera com IDs exclusivos e separadores, mas, nesta comparação, usa somente um veículo e a distância euclidiana original. O experimento não mistura as restrições do problema médico com o TSP da fonte.")

page("10 Sensibilidade e restrições mais rígidas")
picture("outputs/figures/priority_sensitivity.png","Distância e chegada crítica ao variar o coeficiente de prioridade.",height=3.0)
p("A sensibilidade considera coeficientes 0, 0,01, 0,05 e 0,10 em um cenário de 30 entregas e uma semente. A distância e a chegada crítica podem variar sem uma relação monotônica, pois a busca é heurística e tem orçamento finito. Esse gráfico explora a premissa e não estabelece pesos clínicos nem causalidade.")
rows=[]
for name,label in [("sp_stress_capacity","Capacidade 82%"),("sp_stress_range","Autonomia 25 km")]:
    for algo in ("greedy","ortools","genetic"):
        r=result(name,algo,selected if algo=="genetic" else None)
        rows.append([label,algo,fmt(r["objective_mean"]),f"{r['feasible']}/{r['runs']}"])
table(["Variação","Método","J","Viáveis"],rows,[2.2,1.65,1.3,1.35])
p("Na variação de autonomia, OR-Tools não encontrou solução com a estratégia configurada, embora o guloso e o GA tenham encontrado rotas viáveis. Essa falha de busca não prova impossibilidade do problema. O resultado foi mantido, com diagnóstico e sem publicar uma rota inválida.")

page("11 Rotas sobre as ruas de São Paulo")
picture("outputs/maps/sp_evaluation_12.png","Rotas certificadas do cenário de 12 entregas, semente 42.",height=5.15,width=5.7)
p("As linhas acompanham as ruas e retornam ao depósito marcado por uma estrela. As cores dos marcadores representam prioridades e as linhas distinguem os veículos. A geometria é a mesma utilizada nas matrizes. O HTML correspondente permite zoom e seleção de veículos sem carregar mapas ou JavaScript externos.")
p("O traçado demonstra coerência geográfica e respeito ao sentido das vias, mas não valida condições operacionais do dia. Trânsito, bloqueios e restrições de conversão exigiriam um modelo adicional.")

page("12 Frota e chegada das entregas críticas")
picture("outputs/figures/fleet_capacity.png","Carga inicial de cada veículo como percentual de sua capacidade.",height=2.65)
p("A carga inicial inclui todos os pedidos da rota. Os veículos não recebem novas cargas durante o percurso e não realizam coletas. A capacidade é conferida independentemente; utilizar menos de 100% não representa desperdício sem considerar localização, autonomia e prioridades.")
picture("outputs/figures/runtime_critical.png","Tempo de execução e chegada média das entregas críticas.",height=2.45)
p("As chegadas usam minutos após a saída de cada veículo. O tempo de execução é a duração computacional da busca. São medidas distintas. OR-Tools tem cinco segundos; o GA tem limites de fitness e tentativas. A diferença de orçamento impede interpretar estas curvas como uma comparação de eficiência com esforço idêntico.")

page("13 Integração com a LLM e controle factual")
p("Usamos gpt-6-luna pela Responses API com saída estruturada e raciocínio configurado como none. O modelo recebe fatos extraídos de soluções certificadas, a pergunta e as opções autorizadas de organização. Ele seleciona referências, agrupa seções e identifica a sequência de entregas de cada veículo.")
table(["Responsabilidade","Implementação"],[
    ["Métricas e horários","Calculados e certificados em Python"],
    ["Instruções","Veículo, solução e ordem integral de IDs conferidos"],
    ["Relatórios","Fatos diários e agregações semanais previamente calculados"],
    ["Perguntas","Seleção dos fatos pertinentes ou resposta de dados insuficientes"],
    ["Sugestões","Ações condicionais de um vocabulário explícito, sem inventar nova otimização"],
    ["Origem","API real, cache com origem original ou teste identificado"]],[1.45,5.05])
p("Os valores e as frases factuais são renderizados pelo Python. O esquema limita as referências aos fatos existentes e só permite instruções no tipo correspondente. A LLM seleciona e organiza esses fatos, sem escrever livremente novos números ou afirmações clínicas. Essa escolha limita a redação e permite conferência objetiva; a utilidade ainda depende da seleção pertinente ao pedido.")
p("O primeiro prompt apresenta o contrato de seleção de fatos. O segundo explicita o papel de intérprete, o tratamento de conteúdo não confiável, os dados que não podem ser inferidos e a preservação de rotas. Os mesmos 30 casos são usados nos dois prompts.")
p("Perguntas sobre pacientes, temperatura, estoque, trânsito atual, combustível, custo financeiro e prazos clínicos não têm suporte nos dados. A resposta deve reconhecer essa ausência. Uma sugestão de alteração operacional requer nova simulação e certificação; sua menção não comprova melhoria.")

page("14 Avaliação dos prompts e orçamento da LLM")
p("O conjunto fixo contém instruções, relatórios diários e semanais, perguntas respondíveis, oito pedidos de dados ausentes e duas tentativas de indução. Avaliamos status adequado, fatos necessários e preservação da ordem. Toda resposta aceita deve manter consistência factual; a taxa útil nos casos respondíveis deve ser de pelo menos 90%.")
if llm:
    table(["Prompt","Casos","Aceitas","Utilidade","Clareza"],[[v,d["cases"],d["accepted"],fmt(100*d["answerable_useful_fraction"],1)+"%",fmt(d["mean_clarity"],2)] for v,d in llm["prompts"].items()],[1.2,1.0,1.0,1.65,1.65])
    p(f"O resultado do aceite da LLM foi {llm['status']}, com prompt selecionado {llm['selected']}. As respostas rejeitadas permanecem nos registros e reduzem a taxa de utilidade. A origem e o consumo das chamadas são preservados.")
else:
    table(["Evidência","Situação nesta versão"],[["30 casos e dois prompts","Preparados e versionados"],["Chamadas reais","Pendentes de chave específica do projeto"],["Consistência e utilidade reais","Não avaliadas ainda"],["Simulações de falhas e custo","Testadas localmente; não substituem o serviço real"]],[2.15,4.35])
    p("Não há resultado real de qualidade do modelo nesta versão. A ausência da chave não conta como aprovação. O notebook apresenta essa pendência e não identifica respostas simuladas como reais.")
p("A rubrica de clareza considera duplicação, excesso de seções e volume de fatos numa pergunta direta. Avalia uma apresentação controlada; não mede empatia, compreensão clínica ou qualidade de prosa livre. Esse limite deve acompanhar a interpretação das notas.")
initial_path=OUT/"llm/attempts/initial-prompts/summary.json"
if initial_path.exists():
    initial=json.loads(initial_path.read_text(encoding="utf-8"))
    p(f"A rodada inicial falhou: utilidade de {fmt(100*initial['prompts']['v1']['answerable_useful_fraction'],1)}% e {fmt(100*initial['prompts']['v2']['answerable_useful_fraction'],1)}%, com seis dos oito pedidos sem dados corretos em cada prompt. A segunda rodada ainda confundiu uma pergunta sobre recuperação clínica com premissas logísticas. Preservamos ambas e seus custos. Revisamos listas vazias, índices, indução e ausência de informação clínica, além de limitar referências aos fatos existentes. O contrato final é comum aos dois prompts. Como a revisão usou estes casos, as taxas não estimam generalização para perguntas novas.")
p("Antes de cada chamada, a API conta os tokens de entrada e o ledger SQLite reserva o custo máximo, dentro de US$ 1 acumulado. A reserva usa limite conservador de US$ 0,125/milhão de entrada e US$ 0,50/milhão de saída. O uso retornado é reconciliado, e timeouts sem consumo conhecido mantêm a reserva. Não há retries automáticos do SDK.")
if (OUT/"llm/budget-summary.json").exists():
    budget=json.loads((OUT/"llm/budget-summary.json").read_text(encoding="utf-8"));p(f"Compromisso conservador acumulado: US$ {fmt(budget['committed_upper_usd'],6)}, em {len(budget['calls'])} chamadas contabilizadas. O teto inclui falhas, repetições e saídas incompletas.")
p("A chave fica fora do Git e dos notebooks. Respostas de cache são recertificadas contra os fatos atuais e mantêm a origem original. O snapshot JSON do ledger permite reconstruir o banco local conservando o gasto do projeto.")

page("15 Semana simulada")
picture("outputs/figures/weekly_distance.png","Distância diária dos sete cenários fictícios.",height=2.65)
rows=[]
for i in range(1,8):
    data=json.loads((ROOT/f"data/scenarios/sp_day_{i}.json").read_text(encoding="utf-8"))
    solution=json.loads((OUT/f"experiments/runs/sp_day_{i}/genetic_{selected}_42.json").read_text(encoding="utf-8"))
    rows.append([data["day"],solution["metrics"]["deliveries_served"],fmt(solution["metrics"]["total_distance"]/1000),solution["metrics"]["vehicles_used"]])
table(["Dia simulado","Entregas","Distância km","Veículos"],rows,[2,1.25,1.75,1.5])
p(f"O período contém {analysis['weekly_deliveries']} entregas e {fmt(analysis['weekly_distance_km'])} km planejados. A soma dos veículos usados em cada dia representa veículos-dia; não é uma contagem de veículos físicos distintos.")
p("A variação entre dias resulta das localizações, cargas e prioridades sorteadas. Não representa histórico de desempenho de hospitais. O relatório semanal interpreta os totais calculados e preserva as premissas de tempo e simulação.")

page("16 Verificação e reprodução")
p("O certificador independente recalcula visitas, carga, distância, autonomia, horários e métricas sem usar o fitness ou a reparação. O oráculo enumera permutações e divisões de rotas em casos pequenos. Os testes incluem limites exatos, assimetria, entregas no mesmo endereço, pedidos impossíveis, dados inválidos e adulterações de soluções.")
junit=OUT/"validation/pytest.xml"
if junit.exists():p(f"O último conjunto registrado contém {len(ET.parse(junit).findall('.//testcase'))} casos de teste. O XML e o log preservam os resultados e os nomes dos grupos executados.")
p("Na API, verificamos um processo Uvicorn com HTTP real: saúde, Swagger, otimização, consulta persistida e respostas a IDs inexistentes e solicitações inválidas. Os endpoints da LLM dependem de sua própria evidência real. Uma resposta de indisponibilidade não aprova a integração com o provedor.")
p("A conferência geográfica reconstrói matrizes e geometrias a partir do snapshot salvo. Os mapas são abertos no navegador e as figuras estáticas são inspecionadas. Os dois notebooks executam em kernels novos. O relatório é convertido em páginas PNG para conferir tabelas, figuras, quebras de página e legibilidade.")
table(["Comando","Finalidade"],[
    ["tc2 experimentos","Executar ou retomar os experimentos com fingerprints"],
    ["tc2 otimizar sp_evaluation_12 --mapa","Executar um cenário e gerar o mapa"],
    ["tc2 api","Iniciar API local e Swagger"],
    ["tc2 avaliar-llm","Executar ou retomar a avaliação dos dois prompts"],
    ["python scripts/executar_notebooks.py","Executar os dois notebooks integralmente"],
    ["tc2 validar --perfil entrega","Consolidar critérios como aprovado, falhou ou pendente"]],[3.5,3])
p("Os caminhos nos artefatos são relativos ao projeto. A instalação usa ambiente virtual e dependências fixadas após validação. Os arquivos de dados preservam origem, versão e licença. O aceite só é aprovado quando todas as evidências exigidas existem e concordam; pendências externas permanecem pendências.")

page("17 Conclusão e próximos passos")
p("A implementação estende o TSP da base FIAP para múltiplos veículos, prioridades, capacidade e autonomia, com dados geográficos e integração de linguagem fundamentada. A principal contribuição é um fluxo verificável: os algoritmos propõem rotas, o Python certifica e a camada de linguagem interpreta fatos autorizados.")
p("Os experimentos demonstram viabilidade e melhor desempenho que a referência FIAP no TSP compatível, mas não sustentam superioridade geral do genético. Nos problemas maiores, a reparação consome o orçamento e limita a evolução; OR-Tools apresenta custos mais próximos das referências. A análise mantém esses resultados desfavoráveis e distingue falha de busca de prova de impossibilidade.")
p("Como evolução, propomos reparação incremental, inicialização com diversidade viável e avaliação de limites maiores de busca. Essas alterações devem ser medidas novamente nos conjuntos separados. Também seriam necessários trânsito, restrições de conversão, janelas de atendimento e validação dos pesos com especialistas antes de considerar aplicação operacional.")
p("A integração com LLM utiliza apresentação controlada para reduzir erros factuais. Uma futura versão com prosa livre exigiria avaliação semântica mais abrangente, além da conferência de identificadores e valores. A rubrica atual mede seleção e organização de fatos, sem avaliar competência clínica.")
if not llm:p("A entrega integral continua pendente da avaliação real do Luna. Os artefatos locais e os testes disponíveis permitem revisar a implementação, mas não substituem esse requisito. A conclusão da entrega exige resolver a pendência e repetir o aceite completo.")
p("A equipe pode apresentar o problema, os operadores, os dados, as comparações, os mapas e a interpretação em até 14 minutos e 30 segundos. A gravação e a submissão serão realizadas pela equipe, a partir do roteiro e dos comandos incluídos.")

page("18 Referências e localização dos artefatos")
references=[
    ("FIAP  Base de algoritmo genético TSP","https://github.com/FIAP/genetic_algorithm_tsp"),
    ("CVRPLIB  Instâncias e soluções de referência","https://galgos.inf.puc-rio.br/cvrplib/en/instances"),
    ("Espelho CVRPLIB mantido por Thibaut Vidal","https://github.com/vidalt/CVRPLIB"),
    ("OSMnx  Preparação e análise da rede viária","https://osmnx.readthedocs.io/en/stable/user-reference.html"),
    ("OpenStreetMap  Atribuição e licença","https://www.openstreetmap.org/copyright"),
    ("OR-Tools  Vehicle routing problem","https://developers.google.com/optimization/routing/vrp"),
    ("Plotly  Exportação de HTML interativo","https://plotly.com/python/interactive-html-export/"),
    ("OpenAI  Modelo gpt-6-luna","https://developers.openai.com/api/docs/models/gpt-6-luna"),
    ("OpenAI  Contagem de tokens","https://developers.openai.com/api/docs/guides/token-counting"),
    ("OpenAI  Saídas estruturadas na Responses API","https://developers.openai.com/api/docs/guides/structured-outputs?api-mode=responses")]
for title,url in references:p(title+"\n"+url)
p("Fontes consultadas em outubro de 2026. O enunciado IADT Fase 2 e o relatório da fase 1 foram utilizados como referências acadêmicas. Não são incorporados como instruções de execução da aplicação.")
table(["Pasta","Conteúdo"],[["src/rotas_medicas","Domínio, algoritmos, verificação, API, mapas e LLM"],["data","Cenários, matrizes e fontes preservadas"],["outputs/experiments","Soluções, tabelas e parâmetros"],["outputs/figures e maps","Gráficos, figuras e mapas HTML"],["outputs/validation","Evidências e relatório de aceite"],["notebooks e docs","Análises executáveis, arquitetura, plano e roteiro"]],[2.0,4.5])

target=ROOT/"Tech_Challenge_Fase2_Entrega.docx"
doc.core_properties.title="Otimização de rotas médicas com algoritmo genético e LLM"
doc.core_properties.author="Lucas dos Santos Batista; Eduardo Fagner Costalonga Ribeiro; Renan Gonçalves Micheletti"
doc.core_properties.subject="Tech Challenge Fase 2 IADT FIAP"
doc.save(target)
record={"path":target.relative_to(ROOT).as_posix(),"sha256":hashlib.sha256(target.read_bytes()).hexdigest(),
        "analysis_sha256":hashlib.sha256((OUT/"analysis/summary.json").read_bytes()).hexdigest(),"results_sha256":analysis["results_sha256"],
        "required_numeric_text":sorted(set(required_text)),"llm_status":llm["status"] if llm else "pendente","figures":figure_number,
        "layout_reference":"academic structure and figures from phase 1; new document, not an exact visual template"}
(OUT/"analysis/report-data.json").write_text(json.dumps(record,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(target.name,figure_number,"figuras")
